# -*- coding: utf-8 -*-
"""Gera backend/src/client-catalog/client-catalog.data.ts a partir do template.

Uso (a partir da raiz do repositório):
    python docs/tools/gerar-catalogo-dominios.py

Requer: pip install openpyxl

O arquivo gerado é a fonte única dos domínios e notas do cadastro de clientes
(solução S1.2 da ORDEM-02). Formulário, importador, motor de complexidade e
agente de IA leem dele. Nunca editar o .ts à mão: alterar a planilha e rodar
este script.
"""

import datetime
import pathlib
import re
import sys
import unicodedata

import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parents[2]
PLANILHA = RAIZ / "docs" / "Sistema - Base de Cadastro de Clientes.xlsx"
SAIDA = RAIZ / "backend" / "src" / "client-catalog" / "client-catalog.data.ts"
ABA = "CADASTRO DE LISTAS  E COMPLEXID"

FRENTES = ["FISCAL", "CONTABIL", "PESSOAL"]

BLOCO_POR_CADASTRO = {
    "MESTRE": "MESTRE",
    "FISCAL": "FISCAL",
    "CONTÁBIL": "CONTABIL",
    "PESSOAL": "PESSOAL",
}

TIPO = {
    "Texto": "TEXTO",
    "Lista": "LISTA",
    "Mascara": "MASCARA",
    "Switch": "SWITCH",
    "Relação": "RELACAO",
    "Resultado": "RESULTADO",
}

PAPEL = {
    "Cadastro / Controle": "CADASTRO",
    "Elegibilidade": "ELEGIBILIDADE",
    "Complexidade Cliente": "CC",
    "Complexidade Operacional": "CO",
    "Ambos": "AMBOS",
    "Complexidade Cliente - Insumo": "INSUMO",
}

# Opções que significam "esta etapa não existe para este cliente". Saem do
# numerador E do denominador da média (decisão do cliente, reunião de
# 09/09/2026). Não confundir com "Sem integração", que é um estado real e ruim
# e por isso pontua 5.
NAO_SE_APLICA = {"Não se aplica", "Não utiliza controle de ponto"}

# --------------------------------------------------------------------------
# Correção aplicada sobre a planilha
# --------------------------------------------------------------------------
# Campo 44 (PESSOAL / Recebimento documentos) tem as notas de CO preenchidas
# (1,2,3,3,5,3) mas as colunas de classificação ainda dizem "Cadastro /
# Controle" e "CO = Não" — contradição interna do template. Decisão do cliente
# em 09/09/2026: vale a leitura por simetria com Fiscal e Contábil, ou seja,
# o campo PONTUA em CO Pessoal. Fica marcado para confirmação formal.
# Efeito: o divisor do CO Pessoal passa de 4 para 5.
FORCAR_PAPEL = {44: "CO"}


def sem_acento(txt: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", txt)
        if unicodedata.category(c) != "Mn"
    )


def codigo(rotulo: str) -> str:
    slug = sem_acento(rotulo).upper()
    slug = re.sub(r"[^A-Z0-9]+", "_", slug).strip("_")
    return slug or "OPCAO"


def chave_campo(bloco: str, rotulo: str) -> str:
    return f"{bloco}__{codigo(rotulo)}"


def celula(v):
    return str(v).strip() if v is not None else ""


def nota(v):
    if v is None or str(v).strip() == "":
        return None
    try:
        return int(str(v).strip())
    except ValueError:
        return None


def ts(valor: str) -> str:
    return "'" + valor.replace("\\", "\\\\").replace("'", "\\'") + "'"


def main():
    if not PLANILHA.exists():
        sys.exit(f"Planilha não encontrada: {PLANILHA}")

    wb = openpyxl.load_workbook(PLANILHA, data_only=True)
    ws = wb[ABA]

    campos = {}
    ordem = []

    for linha in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=22):
        v = [celula(c.value) for c in linha]
        if not any(v):
            continue
        (num, cad, rotulo, _f, tipo1, _t2, opcao, classe, _cc, _co,
         _k, _l, _m, _n, ccF, ccC, ccP, coF, coC, coP, pilar, regra) = v

        if not num:
            continue
        numero = int(num)
        bloco = BLOCO_POR_CADASTRO.get(cad, cad)
        chave = chave_campo(bloco, rotulo)

        if chave not in campos:
            campos[chave] = {
                "numero": numero,
                "bloco": bloco,
                "rotulo": rotulo,
                "tipo": TIPO.get(tipo1, "TEXTO"),
                "papel": FORCAR_PAPEL.get(numero, PAPEL.get(classe, "CADASTRO")),
                "forcado": numero in FORCAR_PAPEL,
                "pilar": pilar if pilar and pilar != "—" else "",
                "regra": regra,
                "opcoes": [],
                "fronts": set(),
            }
            ordem.append(chave)

        campo = campos[chave]

        notas_cc = dict(zip(FRENTES, (nota(ccF), nota(ccC), nota(ccP))))
        notas_co = dict(zip(FRENTES, (nota(coF), nota(coC), nota(coP))))
        for fr in FRENTES:
            if notas_cc[fr] is not None or notas_co[fr] is not None:
                campo["fronts"].add(fr)

        if opcao and opcao != "-":
            campo["opcoes"].append({
                "codigo": codigo(opcao),
                "rotulo": opcao,
                "cc": {k: n for k, n in notas_cc.items() if n is not None},
                "co": {k: n for k, n in notas_co.items() if n is not None},
                "na": opcao in NAO_SE_APLICA,
                "regra": regra,
            })

    # Campos sem nota nenhuma: a frente vem do bloco.
    for campo in campos.values():
        if not campo["fronts"]:
            if campo["bloco"] == "MESTRE":
                campo["fronts"] = set(FRENTES)
            elif campo["bloco"] in FRENTES:
                campo["fronts"] = {campo["bloco"]}

    hoje = datetime.date.today().strftime("%d/%m/%Y")
    out = []
    w = out.append

    w("// GERADO AUTOMATICAMENTE — não editar à mão.")
    w(f"// Fonte: docs/{PLANILHA.name}")
    w(f"// Gerado em {hoje} por docs/tools/gerar-catalogo-dominios.py")
    w("//")
    w("// Fonte única dos domínios e notas do cadastro de clientes (ORDEM-02,")
    w("// solução S1.2). Consumido pelo formulário, pelo importador, pelo motor")
    w("// de complexidade e pelo agente de IA. Para alterar: mexer na planilha")
    w("// e rodar o gerador novamente.")
    w("")
    w("import {")
    w("  CatalogField,")
    w("  ComplexityFront,")
    w("} from './client-catalog.types';")
    w("")
    w("export const CATALOG_SOURCE = {")
    w(f"  file: {ts(PLANILHA.name)},")
    w(f"  generatedAt: {ts(hoje)},")
    w("};")
    w("")
    w("export const CATALOG_FIELDS: CatalogField[] = [")

    for chave in ordem:
        c = campos[chave]
        fronts = [f for f in FRENTES if f in c["fronts"]]
        w("  {")
        w(f"    key: {ts(chave)},")
        w(f"    number: {c['numero']},")
        w(f"    block: {ts(c['bloco'])},")
        w(f"    label: {ts(c['rotulo'])},")
        w(f"    type: {ts(c['tipo'])},")
        w(f"    role: {ts(c['papel'])},")
        if c["forcado"]:
            w("    // Papel corrigido em relação à planilha: as notas de CO estão")
            w("    // preenchidas mas os marcadores diziam 'Cadastro / Controle'.")
            w("    // Decisão de 09/09/2026, pendente de confirmação formal.")
            w("    roleOverridesTemplate: true,")
        if c["pilar"]:
            w(f"    pillar: {ts(c['pilar'])},")
        w("    fronts: [" + ", ".join(f"{ts(f)} as ComplexityFront" for f in fronts) + "],")
        if c["opcoes"]:
            w("    options: [")
            for o in c["opcoes"]:
                partes = [f"value: {ts(o['codigo'])}", f"label: {ts(o['rotulo'])}"]
                if o["cc"]:
                    partes.append("cc: { " + ", ".join(f"{k}: {v}" for k, v in o["cc"].items()) + " }")
                if o["co"]:
                    partes.append("co: { " + ", ".join(f"{k}: {v}" for k, v in o["co"].items()) + " }")
                if o["na"]:
                    partes.append("notApplicable: true")
                w("      { " + ", ".join(partes) + " },")
            w("    ],")
        w("  },")

    w("];")
    w("")
    w("// Faixas do Total de Vínculos (campo calculado da frente Pessoal).")
    w("// Definidas na planilha, coluna NOTA CC PESSOAL do campo 43.")
    w("export const LINK_COUNT_BANDS = [")
    w("  { max: 10, score: 1 },")
    w("  { max: 50, score: 3 },")
    w("  { max: Infinity, score: 5 },")
    w("];")

    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    SAIDA.write_text("\n".join(out) + "\n", encoding="utf-8")

    pontuam = [c for c in campos.values() if c["papel"] in ("CC", "CO", "AMBOS")]
    print(f"Gerado: {SAIDA.relative_to(RAIZ)}")
    print(f"  {len(campos)} campos, {len(pontuam)} pontuando complexidade")
    for fr in FRENTES:
        cc = [c for c in campos.values() if fr in c["fronts"] and c["papel"] in ("CC", "AMBOS")]
        co = [c for c in campos.values() if fr in c["fronts"] and c["papel"] in ("CO", "AMBOS")]
        print(f"  {fr}: CC={len(cc)} campos, CO={len(co)} campos")


if __name__ == "__main__":
    main()
