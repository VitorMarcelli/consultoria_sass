# -*- coding: utf-8 -*-
"""Regera docs/REFERENCIA-dominios-cadastro-cliente.md a partir do template.

Uso (a partir da raiz do repositório):
    python docs/tools/gerar-referencia-dominios.py

Requer: pip install openpyxl

O documento gerado é a fonte única dos domínios e notas do cadastro de
clientes (solução S1.2 da ORDEM-02). Sempre que a planilha mudar, rode este
script em vez de editar o markdown à mão.
"""

import datetime
import pathlib
import sys

import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parents[2]
PLANILHA = RAIZ / "docs" / "Sistema - Base de Cadastro de Clientes.xlsx"
SAIDA = RAIZ / "docs" / "REFERENCIA-dominios-cadastro-cliente.md"
ABA = "CADASTRO DE LISTAS  E COMPLEXID"

TRAVESSAO = "—"


def celula(valor):
    return str(valor).strip() if valor is not None else ""


def nota(valor):
    return valor or TRAVESSAO


def main():
    if not PLANILHA.exists():
        sys.exit(f"Planilha não encontrada: {PLANILHA}")

    wb = openpyxl.load_workbook(PLANILHA, data_only=True)
    if ABA not in wb.sheetnames:
        sys.exit(f"Aba {ABA!r} não encontrada. Abas disponíveis: {wb.sheetnames}")

    ws = wb[ABA]
    linhas = [
        [celula(c.value) for c in linha]
        for linha in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=22)
    ]

    hoje = datetime.date.today().strftime("%d/%m/%Y")
    out = [
        "# REFERÊNCIA — Domínios e Notas do Cadastro de Clientes",
        "",
        f"**Gerado automaticamente** de `{PLANILHA.name}` em {hoje}.",
        "Não editar à mão: rode `python docs/tools/gerar-referencia-dominios.py`.",
        "",
        "Fonte única para implementar o catálogo de domínios (solução **S1.2** da ORDEM-02).",
        "",
        "## Legenda",
        "",
        "- **CC** = Complexidade do Cliente · **CO** = Complexidade Operacional. Escala 1–5.",
        f"- Notas na ordem **Fiscal / Contábil / Pessoal**; `{TRAVESSAO}` = não pontua naquela frente.",
        "- `N/A` na coluna CO = opção sai do numerador **e** do denominador da média.",
        "- **Atenção:** em *Nota Organização* a escala é invertida — Alta = 1, Baixa = 5.",
        "",
        "---",
        "",
    ]

    atual = None
    for linha in linhas:
        if not any(linha):
            continue
        (num, cad, campo, _f, tipo1, tipo2, opcao, classe, cc, co,
         fiscal, contabil, pessoal, geral,
         cc_f, cc_c, cc_p, co_f, co_c, co_p, pilar, regra) = linha

        chave = (num, cad, campo)
        if chave != atual:
            if atual is not None:
                out.append("")
            atual = chave
            sufixo = f" ({tipo2})" if tipo2 and tipo2 != "-" else ""
            out.append(f"## [{num}] {cad} — {campo}")
            out.append("")
            out.append(f"- **Tipo:** {tipo1}{sufixo}")
            out.append(f"- **Classificação:** {classe} · CC={cc} · CO={co}")
            out.append(
                f"- **Aplica em:** Fiscal={fiscal} · Contábil={contabil} · "
                f"Pessoal={pessoal} · Geral={geral}"
            )
            if pilar and pilar != TRAVESSAO:
                out.append(f"- **Pilar:** {pilar}")
            if opcao and opcao != "-":
                out.append("")
                out.append("| Opção | Nota CC (F/C/P) | Nota CO (F/C/P) | Regra |")
                out.append("|---|---|---|---|")
            elif regra:
                out.append(f"- **Regra:** {regra}")

        if opcao and opcao != "-":
            notas_cc = f"{nota(cc_f)} / {nota(cc_c)} / {nota(cc_p)}"
            notas_co = f"{nota(co_f)} / {nota(co_c)} / {nota(co_p)}"
            texto = (regra or "").replace("|", "\\|")
            out.append(f"| {opcao} | {notas_cc} | {notas_co} | {texto} |")

    SAIDA.write_text("\n".join(out) + "\n", encoding="utf-8")
    print(f"Gerado: {SAIDA.relative_to(RAIZ)} ({len(out)} linhas)")


if __name__ == "__main__":
    main()
