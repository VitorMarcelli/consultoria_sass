# -*- coding: utf-8 -*-
"""Gera o template de carteira em coluna única, a partir do mesmo catálogo.

Uso (a partir da raiz do repositório):
    python docs/tools/gerar-template-carteira.py

Requer: pip install openpyxl

Saída: frontend/public/Template_Carteira_Sevilha.xlsx — o arquivo que o
escritório baixa na tela de importação de clientes.

Por que gerar em vez de manter à mão
------------------------------------
O template é a porta de entrada da carteira. Se um rótulo de coluna ou uma
opção de lista divergir do catálogo (docs/Sistema - Base de Cadastro de
Clientes.xlsx), a importação aceita o arquivo e deixa a resposta em branco —
o cliente entra na carteira sem complexidade e ninguém percebe até o
planejamento sair errado. Gerando template e catálogo da mesma planilha, os
dois não têm como divergir.

Layout
------
Uma aba só, um cliente por linha. As colunas do bloco MESTRE vêm sem prefixo;
as de frente vêm prefixadas — "Fiscal | Nota Volume". O importador reconhece o
prefixo (backend/src/imports/flat-template.ts) e continua aceitando o template
antigo, de quatro abas casadas por CNPJ/CPF.

As listas de opções ficam em colunas ocultas à direita, na mesma aba: é o que
permite ter validação em lista com opções que contêm vírgula (os segmentos) e
com mais de 255 caracteres, sem precisar de uma segunda aba.
"""

import pathlib
import sys
import unicodedata
from collections import OrderedDict

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = pathlib.Path(__file__).resolve().parents[2]
PLANILHA = RAIZ / "docs" / "Sistema - Base de Cadastro de Clientes.xlsx"
SAIDA = RAIZ / "frontend" / "public" / "Template_Carteira_Sevilha.xlsx"
ABA = "CADASTRO DE LISTAS  E COMPLEXID"

LINHAS_PREENCHIVEIS = 500

BLOCOS = ["MESTRE", "FISCAL", "CONTÁBIL", "PESSOAL"]
PREFIXO = {"FISCAL": "Fiscal", "CONTÁBIL": "Contábil", "PESSOAL": "Pessoal"}

# Cor por bloco no cabeçalho. Serve para achar o começo de cada frente numa
# planilha que passa de sessenta colunas.
COR = {
    "MESTRE": "0F766E",
    "FISCAL": "1D4ED8",
    "CONTÁBIL": "7C3AED",
    "PESSOAL": "B45309",
}

PAPEL_TEXTO = {
    "Cadastro / Controle": "informativo, não pontua",
    "Elegibilidade": "define se a frente é avaliada",
    "Complexidade Cliente": "Complexidade do Cliente (CC)",
    "Complexidade Operacional": "Complexidade Operacional (CO)",
    "Ambos": "CC e CO",
    "Complexidade Cliente - Insumo": "alimenta o Total de Vínculos",
}

# Campos calculados pelo sistema: não entram no template porque preenchê-los à
# mão só criaria divergência com o cálculo.
TIPOS_IGNORADOS = {"Resultado"}

STATUS_FRENTE = ["Ativo", "Sem movimento", "Inativo", "Encerrado"]
SIM_NAO = ["Sim", "Não"]


def sem_acento(txt: str) -> str:
    return "".join(
        c
        for c in unicodedata.normalize("NFD", txt)
        if unicodedata.category(c) != "Mn"
    )


def celula(v):
    return str(v).strip() if v is not None else ""


def rotulo_curto(rotulo: str, bloco: str) -> str:
    """Remove o sufixo da frente do rótulo: no template ele já vem no prefixo.

    "Responsável principal - Fiscal" viraria "Fiscal | Responsável principal -
    Fiscal". E é o rótulo curto que o importador procura nas colunas de frente.
    """
    if bloco == "MESTRE":
        return rotulo
    for sufixo in (" - Fiscal", " - Contábil", " - Pessoal", " - Gerais"):
        if rotulo.endswith(sufixo):
            return rotulo[: -len(sufixo)]
    return rotulo


def ler_campos():
    wb = openpyxl.load_workbook(PLANILHA, data_only=True)
    ws = wb[ABA]
    campos = OrderedDict()
    for linha in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=22):
        v = [celula(c.value) for c in linha]
        if not any(v):
            continue
        num, cad, rotulo, _f, tipo1, tipo2, opcao, classe = v[0:8]
        pilar, regra = v[20], v[21]
        if not num:
            continue
        chave = (cad, rotulo)
        if chave not in campos:
            campos[chave] = {
                "num": int(num),
                "bloco": cad,
                "rotulo": rotulo,
                "tipo": tipo1,
                "formato": tipo2,
                "papel": classe,
                "pilar": pilar if pilar and pilar != "—" else "",
                "regra": regra,
                "opcoes": [],
            }
        if opcao and opcao != "-":
            campos[chave]["opcoes"].append(opcao)
    return list(campos.values())


def colunas_do_template(campos):
    """Monta a ordem das colunas: o cliente e depois uma seção por frente."""
    colunas = []
    for bloco in BLOCOS:
        if bloco != "MESTRE":
            # O status da frente não existe no catálogo (que descreve o cadastro,
            # não o ciclo), mas é o que diz ao motor se a frente está parada.
            # Sem ele, "Sem movimento" viraria frente ativa sem respostas.
            colunas.append(
                {
                    "titulo": f"{PREFIXO[bloco]} | Status da frente",
                    "bloco": bloco,
                    "num": None,
                    "tipo": "Lista",
                    "formato": "",
                    "papel": "Elegibilidade",
                    "pilar": "Status / elegibilidade",
                    "regra": (
                        "Situação da frente neste ciclo. 'Sem movimento' e "
                        "'Inativo' tiram a frente do cálculo da complexidade em "
                        "vez de zerá-la."
                    ),
                    "opcoes": STATUS_FRENTE,
                }
            )
        for campo in campos:
            if campo["bloco"] != bloco:
                continue
            if campo["tipo"] in TIPOS_IGNORADOS:
                continue
            curto = rotulo_curto(campo["rotulo"], bloco)
            titulo = curto if bloco == "MESTRE" else f"{PREFIXO[bloco]} | {curto}"
            opcoes = campo["opcoes"]
            if campo["tipo"] == "Switch" and not opcoes:
                opcoes = SIM_NAO
            colunas.append({**campo, "titulo": titulo, "opcoes": opcoes})
    return colunas


def nota_do_cabecalho(coluna) -> str:
    partes = []
    if coluna["num"]:
        partes.append(f"Campo #{coluna['num']} — {coluna['bloco']}")
    else:
        partes.append(f"{coluna['bloco']} — controle do ciclo")
    papel = PAPEL_TEXTO.get(coluna["papel"], coluna["papel"])
    partes.append(f"Entra no cálculo: {papel}")
    if coluna["pilar"]:
        partes.append(f"Pilar: {coluna['pilar']}")
    if coluna["formato"] == "Data":
        # A planilha guarda data como número de dias desde 1899. Pedir o mês
        # escrito evita que a importação receba 46235 onde a pessoa quis dizer
        # agosto de 2026 — o importador converte os dois, mas o texto é o que
        # a pessoa consegue conferir.
        partes.append("Formato: mês no padrão AAAA-MM (ex.: 2026-08)")
    elif coluna["formato"] == "Base de Equipe":
        # A única coluna que aponta para outro cadastro, e por isso a única
        # sem lista de opções: a equipe varia por escritório.
        partes.append(
            "Escreva o nome como está em Estrutura > Equipe. A pessoa precisa "
            "estar cadastrada antes da importação — nome que não bate deixa a "
            "frente sem responsável e aparece no aviso."
        )
    elif coluna["formato"]:
        partes.append(f"Formato: {coluna['formato']}")
    if coluna["opcoes"]:
        partes.append("Opções: " + " · ".join(coluna["opcoes"]))
    regra = coluna["regra"]
    if regra and regra != "—":
        partes.append("")
        partes.append(regra[:700])
    return "\n".join(partes)


def main():
    if not PLANILHA.exists():
        sys.exit(f"Planilha não encontrada: {PLANILHA}")

    campos = ler_campos()
    colunas = colunas_do_template(campos)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Carteira"

    ultima = len(colunas)
    # As listas ficam a duas colunas de distância da última coluna de dados,
    # ocultas: mesma aba, fora do caminho de quem preenche.
    col_listas = ultima + 2

    for i, coluna in enumerate(colunas, start=1):
        letra = get_column_letter(i)
        cel = ws.cell(row=1, column=i, value=coluna["titulo"])
        cel.font = Font(bold=True, color="FFFFFF", size=10)
        cel.fill = PatternFill("solid", fgColor=COR[coluna["bloco"]])
        cel.alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )
        cel.comment = Comment(nota_do_cabecalho(coluna), "Sevilha Performance")
        cel.comment.width = 380
        cel.comment.height = 220

        largura = max(16, min(38, len(coluna["titulo"]) + 4))
        ws.column_dimensions[letra].width = largura

        formato = coluna["formato"]
        if formato == "CNPJ / CPF":
            # Texto, senão o Excel come o zero à esquerda de CPF e o CNPJ vira
            # notação científica.
            estilo = "@"
        elif formato == "Valor":
            estilo = "#,##0.00" if "Honor" in coluna["titulo"] else "0"
        elif formato == "Data":
            estilo = "@"
        else:
            estilo = None
        if estilo:
            for r in range(2, LINHAS_PREENCHIVEIS + 2):
                ws.cell(row=r, column=i).number_format = estilo

        if coluna["opcoes"]:
            letra_lista = get_column_letter(col_listas)
            # A lista começa na linha 2: a linha 1 é o cabeçalho, e quem lê o
            # arquivo monta as colunas a partir dela. Uma opção escrita na
            # linha 1 viraria uma coluna fantasma chamada "Agronegócio".
            for j, opcao in enumerate(coluna["opcoes"], start=2):
                ws.cell(row=j, column=col_listas, value=opcao)
            ws.column_dimensions[letra_lista].hidden = True
            dv = DataValidation(
                type="list",
                formula1=(
                    f"${letra_lista}$2:${letra_lista}${len(coluna['opcoes']) + 1}"
                ),
                allow_blank=True,
                # Aviso em vez de bloqueio: se o escritório tiver um caso que
                # não está na lista, ele precisa conseguir registrar e discutir,
                # não ficar travado na planilha. O importador avisa de novo do
                # lado do sistema.
                errorStyle="warning",
                showErrorMessage=True,
            )
            dv.errorTitle = "Fora da lista"
            dv.error = (
                "Esse valor não está nas opções do campo. Se você continuar, "
                "a importação vai deixar a resposta em branco e avisar."
            )
            ws.add_data_validation(dv)
            dv.add(f"{letra}2:{letra}{LINHAS_PREENCHIVEIS + 1}")
            col_listas += 1

    ws.row_dimensions[1].height = 46
    ws.freeze_panes = "F2"  # mantém CNPJ/CPF e razão social visíveis ao rolar
    ws.auto_filter.ref = f"A1:{get_column_letter(ultima)}1"

    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SAIDA)

    print(f"Gerado: {SAIDA.relative_to(RAIZ)}")
    print(f"  {ultima} colunas em uma aba, {LINHAS_PREENCHIVEIS} linhas com validação")
    for bloco in BLOCOS:
        n = len([c for c in colunas if c["bloco"] == bloco])
        listas = len([c for c in colunas if c["bloco"] == bloco and c["opcoes"]])
        print(f"  {bloco}: {n} colunas ({listas} com lista de opções)")


if __name__ == "__main__":
    main()
